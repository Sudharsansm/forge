import io

from docx import Document
from pypdf import PdfReader

from forge_core import blocks as B

MD = """# Quarterly report

Intro with **bold**, *italic*, `code` and a [link](https://example.com).

## Highlights

- first point
- second **strong** point
  - nested point

1. one
2. two

> a wise quote

| Region | Sales |
|---|---|
| North | 10 |
| South | 20 |

```
print("hi")
```

---

Closing paragraph.
"""


def test_markdown_parse():
    blocks = B.md_to_blocks(MD)
    types = [b["type"] for b in blocks]
    assert types == ["h1", "p", "h2", "li", "li", "li", "oli", "oli", "quote", "table", "code", "hr", "p"]
    intro = blocks[1]["runs"]
    assert any(r.get("b") and r["t"] == "bold" for r in intro)
    assert any(r.get("i") and r["t"] == "italic" for r in intro)
    assert any(r.get("c") and r["t"] == "code" for r in intro)
    assert any(r.get("href") == "https://example.com" for r in intro)
    assert blocks[5]["level"] == 1 and blocks[9]["rows"] == [["Region", "Sales"], ["North", "10"], ["South", "20"]]


def test_markdown_roundtrip_text():
    blocks = B.md_to_blocks(MD)
    again = B.md_to_blocks(B.blocks_to_md(blocks))
    assert [b["type"] for b in again] == [b["type"] for b in blocks]
    assert B.blocks_to_text(again) == B.blocks_to_text(blocks)


def test_snake_case_is_not_italic():
    runs = B.md_to_blocks("use my_var_name here")[0]["runs"]
    assert runs == [{"t": "use my_var_name here"}]


def test_docx_generate_and_read():
    data = B.docx_write(None, B.md_to_blocks(MD), title="Q1")
    back = B.docx_read(data)["blocks"]
    types = [b["type"] for b in back]
    assert types[:3] == ["h1", "p", "h2"] and "table" in types and "li" in types and "oli" in types
    table = next(b for b in back if b["type"] == "table")
    assert table["rows"][1] == ["North", "10"]
    assert any(r.get("b") for r in back[1]["runs"])
    assert Document(io.BytesIO(data)).core_properties.title == "Q1"


def test_docx_in_place_edit_preserves_formatting_and_untouched_paragraphs():
    # an "original" with a custom font that the editor must not destroy
    src = Document()
    p = src.add_paragraph()
    r = p.add_run("Styled text")
    r.font.name = "Georgia"
    r.bold = True
    src.add_paragraph("Second paragraph")
    src.add_heading("A heading", 2)
    buf = io.BytesIO(); src.save(buf)
    original = buf.getvalue()
    blocks = B.docx_read(original)["blocks"]
    assert [b["type"] for b in blocks] == ["p", "p", "h2"]
    blocks[1]["runs"] = [B.run("Second paragraph, edited")]
    blocks.insert(1, {"type": "p", "runs": [B.run("Brand new")]})
    blocks.pop()  # delete the heading
    out = B.docx_write(original, blocks)
    doc = Document(io.BytesIO(out))
    texts = [p.text for p in doc.paragraphs]
    assert texts == ["Styled text", "Brand new", "Second paragraph, edited"]
    assert doc.paragraphs[0].runs[0].font.name == "Georgia" and doc.paragraphs[0].runs[0].bold


def test_docx_table_edit_and_reorder():
    data = B.docx_write(None, B.md_to_blocks("# T\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\nend"))
    blocks = B.docx_read(data)["blocks"]
    tbl = next(b for b in blocks if b["type"] == "table")
    tbl["rows"] = [["a", "b", "c"], ["1", "2", "3"], ["4", "5", "6"]]
    blocks.reverse()
    out = B.docx_read(B.docx_write(data, blocks))["blocks"]
    assert [b["type"] for b in out][0] == "p" and out[-1]["type"] == "h1"
    assert next(b for b in out if b["type"] == "table")["rows"] == tbl["rows"]


def test_pdf_generation():
    pdf = B.blocks_to_pdf(B.md_to_blocks(MD), "Report")
    reader = PdfReader(io.BytesIO(pdf))
    text = "".join(p.extract_text() for p in reader.pages)
    assert len(reader.pages) >= 1 and "Quarterly report" in text and "North" in text and "first point" in text
