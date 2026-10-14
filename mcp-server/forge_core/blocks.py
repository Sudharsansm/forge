"""Editor block model shared by Word (.docx), generated PDF and Markdown conversions.

A block is a JSON object:
    {id, type: h1..h6 | p | li | oli | quote | code | table | image | hr,
     runs:[{t, b?, i?, c?, href?}], level?, rows?:[[str]], src?, ref?, locked?}
`ref` is the index of the element in the original .docx body, so edits can be written back in place and keep
everything the editor does not touch (styles, images, headers, page setup).
"""
from __future__ import annotations

import base64
import copy
import io
import re
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

MONO = {"courier new", "courier", "consolas", "menlo", "monaco", "lucida console", "dejavu sans mono"}
HEADINGS = {f"h{i}" for i in range(1, 7)}
W_P, W_TBL, W_SECT = qn("w:p"), qn("w:tbl"), qn("w:sectPr")
MAX_INLINE_IMAGE = 1_500_000


# ------------------------------------------------------------------------------------------------ runs
def run(t: str, b: bool = False, i: bool = False, c: bool = False, href: str | None = None) -> dict:
    r: dict[str, Any] = {"t": t}
    if b:
        r["b"] = True
    if i:
        r["i"] = True
    if c:
        r["c"] = True
    if href:
        r["href"] = href
    return r


def runs_text(runs: list[dict]) -> str:
    return "".join(r.get("t", "") for r in runs)


def _merge(runs: list[dict]) -> list[dict]:
    out: list[dict] = []
    for r in runs:
        if not r.get("t"):
            continue
        key = lambda x: (bool(x.get("b")), bool(x.get("i")), bool(x.get("c")), x.get("href"))  # noqa: E731
        if out and key(out[-1]) == key(r):
            out[-1] = {**out[-1], "t": out[-1]["t"] + r["t"]}
        else:
            out.append(dict(r))
    return out


_INLINE = re.compile(
    r"(?P<bold>(\*\*|__)(?P<b>.+?)\2)|(?P<it>(?<![\w*])\*(?P<i1>[^*\s][^*]*?)\*(?![\w*])|(?<![\w])_(?P<i2>[^_\s][^_]*?)_(?![\w]))"
    r"|`(?P<code>[^`]+)`|\[(?P<ltxt>[^\]]+)\]\((?P<lurl>[^)\s]+)\)")


def parse_inline(text: str, b: bool = False, i: bool = False) -> list[dict]:
    out: list[dict] = []
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            out.append(run(text[pos:m.start()], b, i))
        if m.group("bold"):
            out += parse_inline(m.group("b"), True, i)
        elif m.group("it"):
            out += parse_inline(m.group("i1") or m.group("i2"), b, True)
        elif m.group("code") is not None:
            out.append(run(m.group("code"), b, i, c=True))
        else:
            out.append(run(m.group("ltxt"), b, i, href=m.group("lurl")))
        pos = m.end()
    if pos < len(text):
        out.append(run(text[pos:], b, i))
    return _merge(out)


def runs_to_md(runs: list[dict]) -> str:
    parts = []
    for r in runs:
        t = r.get("t", "")
        if not t.strip():
            parts.append(t)
            continue
        lead, core, trail = t[: len(t) - len(t.lstrip())], t.strip(), t[len(t.rstrip()):]
        if r.get("c"):
            core = f"`{core}`"
        if r.get("b"):
            core = f"**{core}**"
        if r.get("i"):
            core = f"*{core}*"
        if r.get("href"):
            core = f"[{core}]({r['href']})"
        parts.append(lead + core + trail)
    return "".join(parts)


# ------------------------------------------------------------------------------------------------ markdown
_HR = re.compile(r"^\s{0,3}([-*_])(\s*\1){2,}\s*$")
_LIST = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def md_to_blocks(md: str, plain: bool = False) -> list[dict]:
    lines = md.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[dict] = []
    para: list[str] = []

    def add(b: dict) -> None:
        b["id"] = f"b{len(blocks)}"
        blocks.append(b)

    def flush() -> None:
        if para:
            text = " ".join(s.strip() for s in para)
            add({"type": "p", "runs": [run(text)] if plain else parse_inline(text)})
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        if plain:
            if not line.strip():
                flush()
            else:
                para.append(line)
            i += 1
            continue
        fence = re.match(r"^\s*(```|~~~)", line)
        if fence:
            flush()
            j, code = i + 1, []
            while j < len(lines) and not lines[j].strip().startswith(fence.group(1)):
                code.append(lines[j])
                j += 1
            add({"type": "code", "runs": [run("\n".join(code), c=True)]})
            i = j + 1
            continue
        if not line.strip():
            flush(); i += 1; continue
        h = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", line)
        if h:
            flush()
            add({"type": f"h{len(h.group(1))}", "runs": parse_inline(h.group(2))})
            i += 1; continue
        if _HR.match(line):
            flush(); add({"type": "hr", "runs": []}); i += 1; continue
        if "|" in line and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]) and "-" in lines[i + 1]:
            flush()
            rows = [_cells(line)]
            j = i + 2
            while j < len(lines) and "|" in lines[j] and lines[j].strip():
                rows.append(_cells(lines[j])); j += 1
            width = max(len(r) for r in rows)
            add({"type": "table", "rows": [r + [""] * (width - len(r)) for r in rows]})
            i = j; continue
        if line.lstrip().startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote.append(lines[i].lstrip()[1:].strip()); i += 1
            add({"type": "quote", "runs": parse_inline(" ".join(quote))})
            continue
        m = _LIST.match(line)
        if m:
            flush()
            ordered = m.group(2)[0].isdigit()
            add({"type": "oli" if ordered else "li", "level": min(len(m.group(1).replace("\t", "  ")) // 2, 4), "runs": parse_inline(m.group(3))})
            i += 1; continue
        img = re.match(r"^\s*!\[([^\]]*)\]\(([^)\s]+)\)\s*$", line)
        if img:
            flush()
            if img.group(2).startswith("data:image/"):
                add({"type": "image", "src": img.group(2), "runs": [run(img.group(1))]})
            else:
                add({"type": "p", "runs": [run(f"[image: {img.group(1) or img.group(2)}]", i=True)]})
            i += 1; continue
        para.append(line)
        i += 1
    flush()
    return blocks


def blocks_to_md(blocks: list[dict]) -> str:
    out: list[str] = []
    n = 0
    for b in blocks:
        t = b.get("type", "p")
        if t != "oli":
            n = 0
        if t in HEADINGS:
            out.append("#" * int(t[1]) + " " + runs_to_md(b.get("runs", [])))
        elif t == "li":
            out.append("  " * int(b.get("level", 0)) + "- " + runs_to_md(b.get("runs", [])))
        elif t == "oli":
            n += 1
            out.append("  " * int(b.get("level", 0)) + f"{n}. " + runs_to_md(b.get("runs", [])))
        elif t == "quote":
            out.append("> " + runs_to_md(b.get("runs", [])))
        elif t == "code":
            out.append("```\n" + runs_text(b.get("runs", [])) + "\n```")
        elif t == "hr":
            out.append("---")
        elif t == "image":
            out.append(f"![{runs_text(b.get('runs', [])) or 'image'}]()")
        elif t == "table":
            rows = b.get("rows") or [[""]]
            esc = lambda c: str(c).replace("|", "\\|").replace("\n", " ")  # noqa: E731
            out.append("\n".join(["| " + " | ".join(esc(c) for c in rows[0]) + " |", "|" + "---|" * len(rows[0])]
                                 + ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows[1:]]))
        else:
            out.append(runs_to_md(b.get("runs", [])))
        if t not in ("li", "oli"):
            out.append("")
        elif out and False:
            pass
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"


def blocks_to_text(blocks: list[dict], with_ids: bool = False) -> str:
    out = []
    for b in blocks:
        t = b.get("type", "p")
        body = ("\n".join("\t".join(r) for r in b.get("rows", [])) if t == "table"
                else "[image]" if t == "image" else "---" if t == "hr" else runs_text(b.get("runs", [])))
        prefix = f"[{b.get('id', '')}:{t}] " if with_ids else ""
        out.append(prefix + body)
    return "\n".join(out)


# ------------------------------------------------------------------------------------------------ docx: read
def _style_type(p: Paragraph) -> tuple[str, int]:
    name = (p.style.name if p.style is not None else "") or ""
    low = name.lower()
    m = re.match(r"heading (\d)", low)
    if m:
        return f"h{min(6, int(m.group(1)))}", 0
    if low == "title":
        return "h1", 0
    if low == "subtitle":
        return "h2", 0
    m = re.match(r"list (bullet|number)(?: (\d))?", low)
    if m:
        return ("li" if m.group(1) == "bullet" else "oli"), max(0, int(m.group(2) or 1) - 1)
    if "quote" in low:
        return "quote", 0
    ppr = p._p.pPr
    if ppr is not None and ppr.numPr is not None:
        return "li", int(ppr.numPr.ilvl.val) if ppr.numPr.ilvl is not None else 0
    return "p", 0


def _read_runs(p: Paragraph) -> list[dict]:
    runs: list[dict] = []
    for item in p.iter_inner_content():
        if isinstance(item, Run):
            font = (item.font.name or "").lower()
            runs.append(run(item.text, bool(item.bold), bool(item.italic), font in MONO))
        else:  # Hyperlink
            runs.append(run(item.text, href=item.url or None))
    return _merge(runs)


def _paragraph_image(p: Paragraph, doc) -> str | None:
    for blip in p._p.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid and rid in doc.part.related_parts:
            part = doc.part.related_parts[rid]
            blob = part.blob
            if len(blob) <= MAX_INLINE_IMAGE and part.content_type.startswith("image/") and "tiff" not in part.content_type:
                return f"data:{part.content_type};base64,{base64.b64encode(blob).decode()}"
            return ""  # present but not shown
    return None


def docx_read(data: bytes) -> dict:
    doc = Document(io.BytesIO(data))
    blocks: list[dict] = []
    for idx, el in enumerate(el for el in doc.element.body.iterchildren() if el.tag in (W_P, W_TBL)):
        bid = f"b{idx}"
        if el.tag == W_TBL:
            t = Table(el, doc._body)
            blocks.append({"id": bid, "ref": idx, "type": "table", "rows": [[c.text for c in row.cells] for row in t.rows]})
            continue
        p = Paragraph(el, doc._body)
        has_draw = bool(list(el.iter(qn("w:drawing"))) or list(el.iter(qn("w:pict"))))
        runs = _read_runs(p)
        if has_draw:
            src = _paragraph_image(p, doc)
            if not runs_text(runs).strip():
                blocks.append({"id": bid, "ref": idx, "type": "image", "src": src or "", "runs": [], "locked": True})
                continue
            kind, level = _style_type(p)
            blocks.append({"id": bid, "ref": idx, "type": kind, "level": level, "runs": runs, "locked": True})
            continue
        kind, level = _style_type(p)
        if kind == "p" and runs and all(r.get("c") for r in runs):
            kind = "code"
        if not runs and kind == "p" and not p.text:
            runs = []
        blocks.append({"id": bid, "ref": idx, "type": kind, "level": level, "runs": runs})
    core = doc.core_properties
    return {"blocks": blocks, "title": core.title or "", "author": core.author or ""}


# ------------------------------------------------------------------------------------------------ docx: write
_STYLE_FOR = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "h4": "Heading 4", "h5": "Heading 5", "h6": "Heading 6",
              "li": "List Bullet", "oli": "List Number", "quote": "Quote", "p": "Normal", "code": "Normal", "hr": "Normal"}


def _set_style(doc, p: Paragraph, btype: str, level: int = 0) -> None:
    name = _STYLE_FOR.get(btype, "Normal")
    if btype in ("li", "oli") and level > 0:
        name = f"{name} {min(level + 1, 3)}"
    for candidate in (name, _STYLE_FOR.get(btype, "Normal"), "Normal"):
        try:
            p.style = doc.styles[candidate]
            return
        except KeyError:
            continue


def _clear_runs(p: Paragraph) -> OxmlElement | None:
    """Remove text runs and hyperlinks; keep the first run's properties so the font and size survive."""
    base = None
    for child in list(p._p):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            if base is None:
                r = child if child.tag == qn("w:r") else child.find(qn("w:r"))
                if r is not None and r.find(qn("w:rPr")) is not None:
                    base = copy.deepcopy(r.find(qn("w:rPr")))
            p._p.remove(child)
    return base


def _add_runs(doc, p: Paragraph, runs: list[dict], base_rpr, code: bool = False) -> None:
    for r in runs:
        text = r.get("t", "")
        if not text:
            continue
        if r.get("href"):
            rid = doc.part.relate_to(r["href"], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), rid)
            wr = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            style = OxmlElement("w:rStyle")
            style.set(qn("w:val"), "Hyperlink")
            rpr.append(style)
            color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rpr.append(color)
            u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
            wr.append(rpr)
            wt = OxmlElement("w:t"); wt.text = text; wt.set(qn("xml:space"), "preserve"); wr.append(wt)
            link.append(wr)
            p._p.append(link)
            continue
        run_ = p.add_run()
        # honour line breaks / tabs inside a run
        parts = re.split(r"(\n|\t)", text)
        for part in parts:
            if part == "\n":
                run_.add_break()
            elif part == "\t":
                run_.add_tab()
            elif part:
                run_.add_text(part)
        if base_rpr is not None:
            old = run_._r.find(qn("w:rPr"))
            if old is not None:
                run_._r.remove(old)
            run_._r.insert(0, copy.deepcopy(base_rpr))
        run_.bold = True if r.get("b") else (False if base_rpr is not None else None)
        run_.italic = True if r.get("i") else (False if base_rpr is not None else None)
        if r.get("c") or code:
            run_.font.name = "Courier New"
            run_.font.size = Pt(10)


def _hr(p: Paragraph) -> None:
    ppr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(k), v)
    border.append(bottom)
    ppr.append(border)


def _fill_table(doc, table: Table, rows: list[list[str]]) -> None:
    rows = rows or [[""]]
    width = max(len(r) for r in rows)
    while len(table.rows) < len(rows):
        table.add_row()
    for extra in range(len(table.rows) - len(rows)):
        tr = table.rows[len(table.rows) - 1]._tr
        tr.getparent().remove(tr)
    while len(table.columns) < width:
        table.add_column(Pt(72))
    for ri, row in enumerate(rows):
        cells = table.rows[ri].cells
        for ci in range(min(width, len(cells))):
            want = str(row[ci]) if ci < len(row) else ""
            if cells[ci].text != want:
                cells[ci].text = want


def _image_to_stream(src: str):
    m = re.match(r"data:image/(png|jpe?g|gif|bmp);base64,(.+)$", src or "", re.S)
    return io.BytesIO(base64.b64decode(m.group(2))) if m else None


def docx_write(original: bytes | None, blocks: list[dict], title: str | None = None) -> bytes:
    doc = Document(io.BytesIO(original)) if original else Document()
    body = doc.element.body
    orig_items = [el for el in body.iterchildren() if el.tag in (W_P, W_TBL)]
    if original is None:  # drop the empty paragraph a new document may start with
        for el in orig_items:
            body.remove(el)
        orig_items = []
    used: set[int] = set()
    final: list[Any] = []
    for b in blocks:
        btype = b.get("type", "p")
        runs = b.get("runs") or []
        ref = b.get("ref")
        el = orig_items[ref] if isinstance(ref, int) and 0 <= ref < len(orig_items) and ref not in used else None
        if el is not None and ((btype == "table") != (el.tag == W_TBL)):
            el = None
        if el is not None:
            used.add(ref)
            if el.tag == W_TBL:
                _fill_table(doc, Table(el, doc._body), b.get("rows") or [[""]])
                final.append(el)
                continue
            p = Paragraph(el, doc._body)
            if b.get("locked"):
                final.append(el)
                continue
            old_type, old_level = _style_type(p)
            if old_type != btype or old_level != int(b.get("level", 0) or 0):
                if not (old_type == "code" and btype == "code"):
                    _set_style(doc, p, btype, int(b.get("level", 0) or 0))
            if _read_runs(p) != _merge(runs):
                base = _clear_runs(p)
                _add_runs(doc, p, runs, base, code=btype == "code")
            final.append(el)
            continue
        # a new element
        if btype == "table":
            rows = b.get("rows") or [[""]]
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            try:
                t.style = doc.styles["Table Grid"]
            except KeyError:
                pass
            _fill_table(doc, t, rows)
            final.append(t._tbl)
            continue
        if btype == "image":
            stream = _image_to_stream(b.get("src", ""))
            p = doc.add_paragraph()
            if stream is not None:
                p.add_run().add_picture(stream, width=Pt(380))
            final.append(p._p)
            continue
        p = doc.add_paragraph()
        _set_style(doc, p, btype, int(b.get("level", 0) or 0))
        if btype == "hr":
            _hr(p)
        else:
            _add_runs(doc, p, runs, None, code=btype == "code")
        final.append(p._p)
    for el in orig_items:
        if el.getparent() is not None:
            body.remove(el)
    for el in final:
        if el.getparent() is not None:
            body.remove(el)
    sect = body.find(W_SECT)
    for el in final:
        if sect is not None:
            sect.addprevious(el)
        else:
            body.append(el)
    if title:
        doc.core_properties.title = title
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ------------------------------------------------------------------------------------------------ pdf
def _pdf_font() -> tuple[str, str, str, str]:
    """(regular, bold, italic, mono) font names; DejaVu when installed so accented and non-Latin text renders."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os
    candidates = ["/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "C:/Windows/Fonts", "/Library/Fonts", "/System/Library/Fonts/Supplemental"]
    for d in candidates:
        reg = os.path.join(d, "DejaVuSans.ttf")
        if os.path.isfile(reg):
            try:
                pdfmetrics.registerFont(TTFont("ForgeSans", reg))
                bold = os.path.join(d, "DejaVuSans-Bold.ttf")
                ital = os.path.join(d, "DejaVuSans-Oblique.ttf")
                mono = os.path.join(d, "DejaVuSansMono.ttf")
                pdfmetrics.registerFont(TTFont("ForgeSans-Bold", bold if os.path.isfile(bold) else reg))
                pdfmetrics.registerFont(TTFont("ForgeSans-Italic", ital if os.path.isfile(ital) else reg))
                pdfmetrics.registerFont(TTFont("ForgeMono", mono if os.path.isfile(mono) else reg))
                return "ForgeSans", "ForgeSans-Bold", "ForgeSans-Italic", "ForgeMono"
            except Exception:  # noqa: BLE001 - fall back to the built-in fonts
                break
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Courier"


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _runs_markup(runs: list[dict], mono: str) -> str:
    out = []
    for r in runs:
        t = _esc(r.get("t", "")).replace("\n", "<br/>")
        if r.get("c"):
            t = f'<font name="{mono}" size="9">{t}</font>'
        if r.get("b"):
            t = f"<b>{t}</b>"
        if r.get("i"):
            t = f"<i>{t}</i>"
        if r.get("href"):
            t = f'<a href="{_esc(r["href"])}" color="#2563eb">{t}</a>'
        out.append(t)
    return "".join(out)


def blocks_to_pdf(blocks: list[dict], title: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, Image, ListFlowable, ListItem, Paragraph as RLParagraph, Preformatted, SimpleDocTemplate, Spacer, Table as RLTable, TableStyle
    from PIL import Image as PILImage

    reg, bold, ital, mono = _pdf_font()
    body = ParagraphStyle("body", fontName=reg, fontSize=10.5, leading=15, spaceAfter=6, alignment=TA_LEFT)
    heads = {n: ParagraphStyle(f"h{n}", parent=body, fontName=bold, fontSize=s, leading=s + 5, spaceBefore=12 if n > 1 else 4, spaceAfter=6)
             for n, s in zip(range(1, 7), (22, 17, 14, 12, 11, 10.5))}
    quote = ParagraphStyle("quote", parent=body, fontName=ital, leftIndent=14, textColor=colors.HexColor("#555555"))
    story: list[Any] = []
    pending: list[tuple[str, str, int]] = []

    def flush_list() -> None:
        if not pending:
            return
        kind = pending[0][0]
        items = [ListItem(RLParagraph(m, body), leftIndent=14 + 14 * lvl) for _, m, lvl in pending]
        story.append(ListFlowable(items, bulletType="1" if kind == "oli" else "bullet", start=None if kind != "oli" else 1,
                                  bulletFontName=reg, leftIndent=14))
        pending.clear()

    for b in blocks:
        t = b.get("type", "p")
        if t in ("li", "oli"):
            if pending and pending[0][0] != t:
                flush_list()
            pending.append((t, _runs_markup(b.get("runs", []), mono), int(b.get("level", 0) or 0)))
            continue
        flush_list()
        if t in HEADINGS:
            story.append(RLParagraph(_runs_markup(b.get("runs", []), mono), heads[int(t[1])]))
        elif t == "quote":
            story.append(RLParagraph(_runs_markup(b.get("runs", []), mono), quote))
        elif t == "code":
            story.append(Preformatted(runs_text(b.get("runs", [])), ParagraphStyle("code", fontName=mono, fontSize=8.5, leading=11,
                                                                                    backColor=colors.HexColor("#f4f4f2"), borderPadding=6, spaceAfter=10)))
        elif t == "hr":
            story.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc"), spaceBefore=6, spaceAfter=10))
        elif t == "table":
            rows = b.get("rows") or [[""]]
            data = [[RLParagraph(_esc(str(c)), ParagraphStyle("cell", parent=body, fontSize=9, leading=12, spaceAfter=0,
                                                              fontName=bold if ri == 0 else reg)) for c in r] for ri, r in enumerate(rows)]
            tbl = RLTable(data, repeatRows=1)
            tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
                                     ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0ee")),
                                     ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
            story += [tbl, Spacer(1, 10)]
        elif t == "image":
            stream = _image_to_stream(b.get("src", ""))
            if stream is not None:
                with PILImage.open(stream) as im:
                    w, h = im.size
                stream.seek(0)
                width = min(15 * cm, w * 0.75)
                story += [Image(stream, width=width, height=width * h / w), Spacer(1, 8)]
        else:
            markup = _runs_markup(b.get("runs", []), mono)
            story.append(RLParagraph(markup if markup else "&nbsp;", body))
    flush_list()
    if not story:
        story.append(RLParagraph("&nbsp;", body))
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                      title=title or "Document", author="Forge").build(story)
    return out.getvalue()


def ensure_ids(blocks: list[dict]) -> list[dict]:
    for n, b in enumerate(blocks):
        b.setdefault("id", f"b{n}")
    return blocks
